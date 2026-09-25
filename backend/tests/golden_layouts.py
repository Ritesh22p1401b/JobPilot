"""Programmatically generated 'golden' resume layouts covering known parser risk patterns."""

from __future__ import annotations

import io

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

CONTACT = "Asha Verma | asha.verma@example.com | +91 99887 76655 | Pune, India"
EXPERIENCE = [
    ("Backend Engineer | Nimbus Tech | Pune", "Mar 2022 – Present",
     ["Built REST APIs with FastAPI and PostgreSQL for billing services.",
      "Containerized microservices with Docker and deployed them on AWS."]),
    ("Software Engineer Intern | Orbit Labs", "Jun 2021 – Dec 2021", ["Wrote Python ETL jobs using Pandas."]),
]
SKILLS = "Python, FastAPI, PostgreSQL, Docker, AWS, Pandas"
EDU = "B.E. in Information Technology | Savitribai Phule Pune University | 2017 – 2021"


def _base_docx() -> Document:
    d = Document()
    d.add_paragraph("Asha Verma")
    return d


def _body(d: Document, contact: bool = True) -> None:
    if contact:
        d.add_paragraph(CONTACT)
    d.add_paragraph("EXPERIENCE")
    for head, dates, bullets in EXPERIENCE:
        d.add_paragraph(head)
        d.add_paragraph(dates)
        for b in bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_paragraph("SKILLS")
    d.add_paragraph(SKILLS)
    d.add_paragraph("EDUCATION")
    d.add_paragraph(EDU)


def _save(d: Document) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def single_column_docx() -> bytes:
    d = _base_docx()
    _body(d)
    return _save(d)


def table_docx() -> bytes:
    d = _base_docx()
    d.add_paragraph(CONTACT)
    d.add_paragraph("EXPERIENCE")
    t = d.add_table(rows=len(EXPERIENCE), cols=2)
    for i, (head, dates, bullets) in enumerate(EXPERIENCE):
        t.cell(i, 0).text = f"{head}\n{dates}"
        t.cell(i, 1).text = "\n".join(bullets)
    d.add_paragraph("SKILLS")
    d.add_paragraph(SKILLS)
    d.add_paragraph("EDUCATION")
    d.add_paragraph(EDU)
    return _save(d)


def header_contact_docx() -> bytes:
    d = _base_docx()
    d.sections[0].header.paragraphs[0].text = CONTACT
    _body(d, contact=False)
    return _save(d)


def text_box_docx() -> bytes:
    d = _base_docx()
    _body(d)
    p = d.add_paragraph()
    run = p.add_run()
    textbox = parse_xml(
        f'<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><v:shape style="width:200pt;height:40pt">'
        f'<v:textbox><w:txbxContent><w:p><w:r><w:t>Certified Kubernetes Administrator</w:t></w:r></w:p>'
        f"</w:txbxContent></v:textbox></v:shape></w:pict>"
    )
    run._r.append(textbox)
    return _save(d)


def icons_docx() -> bytes:
    d = _base_docx()
    d.add_paragraph(" asha.verma@example.com   +91 99887 76655   Pune, India")
    _body(d, contact=False)
    for _ in range(3):
        d.add_paragraph(" Python   Docker")
    return _save(d)


def image_docx() -> bytes:
    from PIL import Image

    img = io.BytesIO()
    Image.new("RGB", (40, 40), (30, 90, 200)).save(img, format="PNG")
    img.seek(0)
    d = _base_docx()
    d.add_picture(img)
    _body(d)
    return _save(d)


def columns_docx() -> bytes:
    d = _base_docx()
    _body(d)
    sect = d.sections[0]._sectPr
    cols = sect.find(qn("w:cols"))
    if cols is None:
        cols = parse_xml(f'<w:cols {nsdecls("w")} w:num="2"/>')
        sect.append(cols)
    else:
        cols.set(qn("w:num"), "2")
    return _save(d)


def two_column_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import BaseDocTemplate, Frame, FrameBreak, PageTemplate, Paragraph

    buf = io.BytesIO()
    w, h = A4
    doc = BaseDocTemplate(buf, pagesize=A4)
    left = Frame(30, 40, w * 0.32, h - 80, id="left")
    right = Frame(40 + w * 0.32, 40, w * 0.6, h - 80, id="right")
    doc.addPageTemplates([PageTemplate(frames=[left, right])])
    st = getSampleStyleSheet()["Normal"]
    story = [Paragraph("Asha Verma", st), Paragraph("asha.verma@example.com", st), Paragraph("+91 99887 76655", st),
             Paragraph("SKILLS", st)] + [Paragraph(s.strip(), st) for s in SKILLS.split(",")] + [FrameBreak()]
    story.append(Paragraph("EXPERIENCE", st))
    for head, dates, bullets in EXPERIENCE:
        story += [Paragraph(head, st), Paragraph(dates, st)] + [Paragraph("• " + b, st) for b in bullets]
    story += [Paragraph("EDUCATION", st), Paragraph(EDU, st)] + [Paragraph(" ".join(EXPERIENCE[0][2]), st) for _ in range(6)]
    doc.build(story)
    return buf.getvalue()


GOLDEN = {
    "single_column": ("docx", single_column_docx, set()),
    "tables": ("docx", table_docx, {"TABLES"}),
    "header_contact": ("docx", header_contact_docx, {"HEADER_CONTACT"}),
    "text_box": ("docx", text_box_docx, {"TEXT_BOXES"}),
    "icons": ("docx", icons_docx, {"UNUSUAL_SYMBOLS"}),
    "image": ("docx", image_docx, {"IMAGES"}),
    "columns_docx": ("docx", columns_docx, {"MULTI_COLUMN"}),
    "two_column_pdf": ("pdf", two_column_pdf, {"MULTI_COLUMN"}),
}
